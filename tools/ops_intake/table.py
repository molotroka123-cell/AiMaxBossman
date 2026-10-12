"""Tables: one row per unique request, a duplicates table, needs_review, skipped. CSV always; XLSX if openpyxl.

CSV is UTF-8 with BOM (Excel opens Cyrillic correctly). Cells that a spreadsheet could execute as a formula
(leading ``=`` TAB CR, an ``@`` that is not a plain @handle, or a ``+``/``-`` that is not a number) are prefixed with ``'`` in CSV and written as explicit
strings in XLSX - request text comes from untrusted senders.
"""
from __future__ import annotations

import csv
import re
from decimal import Decimal
from pathlib import Path
from typing import Optional

from .dedupe import DedupeResult
from .extract import REVIEW_FLAGS

COLUMNS = [
    "row_id", "first_src_id", "ts_utc", "source", "sender", "lead_id", "name", "phone", "phone_raw", "phone_alt",
    "email", "email_raw", "handle", "service", "service_candidates", "amount", "currency", "amount_raw",
    "amount_candidates", "deadline_date", "deadline_raw", "deadline_basis", "language", "needs_review",
    "review_reasons", "flags", "duplicates_count", "merged_src_ids", "filled_from", "possible_duplicate_of",
    "answered", "answered_basis", "text",
]
DUPLICATE_COLUMNS = [
    "dup_src_id", "kept_src_id", "kept_row_id", "rule", "confidence", "delta_min", "window_verified", "similarity",
    "dup_ts_utc", "dup_sender", "name", "phone", "email", "service", "amount", "currency", "new_info", "text",
]
SKIPPED_COLUMNS = ["src_id", "category", "ts_utc", "sender", "text_or_reason"]
REQUIRE_CHOICES = ("name", "phone", "email", "service", "amount", "deadline", "contact")

_CTRL_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
_NUMERIC_LIKE_RE = re.compile(r"^[+-]?[\d\s().,-]+$")
_HANDLE_RE = re.compile(r"^@[A-Za-z][A-Za-z0-9_]{4,31}$")      # a plain Telegram @handle cannot be a formula


def safe_cell(value) -> str:
    s = "" if value is None else str(value)
    s = _CTRL_RE.sub("", s)
    if s and (s[0] in "=\t\r" or (s[0] == "@" and not _HANDLE_RE.match(s))
              or (s[0] in "+-" and not _NUMERIC_LIKE_RE.match(s))):
        s = "'" + s
    return s


def _one_line(s: str, limit: int = 1000) -> str:
    s = re.sub(r"\s+", " ", s or "").strip()
    return s if len(s) <= limit else s[: limit - 1] + "…"


def _amount_str(a: Optional[Decimal]) -> str:
    if a is None:
        return ""
    return format(a.normalize(), "f") if a == a.to_integral() else str(a.normalize())


def _ts(t) -> str:
    return t.strftime("%Y-%m-%d %H:%M:%S") if t else ""


def _answered(v) -> str:
    return {True: "yes", False: "no", None: "unknown"}[v]


def review_reasons(flags: list, require: tuple) -> list:
    reasons = [f for f in flags if f in REVIEW_FLAGS]
    for field_name in require:
        key = f"{field_name}_missing"
        if key in flags and key not in reasons:
            reasons.append(key)
    return reasons


def build_rows(result: DedupeResult, require: tuple = ("service", "contact")) -> list:
    bad = [r for r in require if r not in REQUIRE_CHOICES]
    if bad:
        raise ValueError(f"unknown --require field(s) {bad}; choose from {list(REQUIRE_CHOICES)}")
    ids = {id(c): f"R{n:04d}" for n, c in enumerate(result.clusters, start=1)}
    rows: list = []
    for c in result.clusters:
        e = c.row
        reasons = review_reasons(e.flags, tuple(require))
        rows.append({
            "row_id": ids[id(c)], "first_src_id": c.first_src_id, "ts_utc": _ts(e.ts_utc), "source": e.source,
            "sender": e.sender, "lead_id": e.lead_id, "name": e.name, "phone": e.phone, "phone_raw": e.phone_raw,
            "phone_alt": "; ".join(e.phone_alt), "email": e.email, "email_raw": e.email_raw, "handle": e.handle,
            "service": e.service, "service_candidates": "; ".join(e.service_candidates),
            "amount": _amount_str(e.amount), "currency": e.currency, "amount_raw": e.amount_raw,
            "amount_candidates": "; ".join(e.amount_candidates), "deadline_date": e.deadline_date,
            "deadline_raw": e.deadline_raw, "deadline_basis": e.deadline_basis, "language": e.language,
            "needs_review": "yes" if reasons else "no", "review_reasons": "; ".join(reasons),
            "flags": "; ".join(e.flags), "duplicates_count": len(c.links),
            "merged_src_ids": "; ".join(l.dup.src_id for l in c.links), "filled_from": "; ".join(c.filled),
            "possible_duplicate_of": "; ".join(ids[id(p["other"])] for p in c.possible),
            "answered": _answered(e.answered), "answered_basis": e.answered_basis, "text": _one_line(c.members[0].text),
        })
    return rows


def build_duplicate_rows(result: DedupeResult, rows: list) -> list:
    row_of = {r["first_src_id"]: r["row_id"] for r in rows}
    out = []
    for l in result.links:
        d = l.dup
        out.append({
            "dup_src_id": d.src_id, "kept_src_id": l.kept_src_id, "kept_row_id": row_of.get(l.kept_src_id, ""),
            "rule": l.rule, "confidence": l.confidence, "delta_min": "" if l.delta_min is None else l.delta_min,
            "window_verified": "yes" if l.window_verified else "no",
            "similarity": "" if l.similarity is None else round(l.similarity, 3), "dup_ts_utc": _ts(d.ts_utc),
            "dup_sender": d.sender, "name": d.name, "phone": d.phone, "email": d.email, "service": d.service,
            "amount": _amount_str(d.amount), "currency": d.currency, "new_info": "; ".join(l.new_info),
            "text": _one_line(d.text),
        })
    return out


def build_skipped_rows(skipped: list, unparsed: list) -> list:
    out = []
    for e in skipped:
        out.append({"src_id": e.src_id, "category": e.skip_reason, "ts_utc": _ts(e.ts_utc), "sender": e.sender,
                    "text_or_reason": _one_line(e.text, 200)})
    for u in unparsed:
        out.append({"src_id": u.get("ref", ""), "category": "unparsed", "ts_utc": "", "sender": "",
                    "text_or_reason": u.get("reason", "")})
    return out


def write_csv(path, columns: list, rows: list) -> None:
    with open(path, "w", encoding="utf-8-sig", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(columns)
        for r in rows:
            w.writerow([safe_cell(r.get(c, "")) for c in columns])


def _load_openpyxl():
    import openpyxl
    return openpyxl


def write_xlsx(path, sheets: dict) -> dict:
    """``sheets``: name -> (columns, rows). Returns {"written": bool, "reason": str, "sheets": [...]}.
    Degrades honestly: without openpyxl nothing is written and the reason says so."""
    try:
        openpyxl = _load_openpyxl()
    except ImportError:
        return {"written": False, "reason": "openpyxl is not importable; the CSV files are the table", "sheets": []}
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, (columns, rows) in sheets.items():
        ws = wb.create_sheet(name[:31])
        for ci, col in enumerate(columns, start=1):
            ws.cell(row=1, column=ci, value=col)
        for ri, r in enumerate(rows, start=2):
            for ci, col in enumerate(columns, start=1):
                v = r.get(col, "")
                cell = ws.cell(row=ri, column=ci)
                if col == "amount" and v not in ("", None):
                    cell.value = float(Decimal(str(v)))
                    continue
                if isinstance(v, (int, float)) and not isinstance(v, bool):
                    cell.value = v
                    continue
                cell.value = _CTRL_RE.sub("", "" if v is None else str(v))
                cell.data_type = "s"        # never a formula, even when the text starts with "="
        ws.freeze_panes = "A2"
        if columns and rows:
            ws.auto_filter.ref = f"A1:{openpyxl.utils.get_column_letter(len(columns))}{len(rows) + 1}"
        for ci, col in enumerate(columns, start=1):
            width = max([len(col)] + [min(len(str(r.get(col, ""))), 60) for r in rows[:200]])
            ws.column_dimensions[openpyxl.utils.get_column_letter(ci)].width = min(max(width + 2, 8), 62)
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return {"written": True, "reason": "", "sheets": list(sheets), "openpyxl": getattr(openpyxl, "__version__", "")}
