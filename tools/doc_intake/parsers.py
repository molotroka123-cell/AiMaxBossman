"""Local-only document parsers. Nothing here opens a socket.

Open-source libraries first, feature-detected at call time:

* PDF  — Docling (MIT) only when explicitly enabled AND installed (tables as
  Markdown, repo adapter ``bcc.oss.docling`` = native pipeline, no model
  download); otherwise pypdf (BSD) page by page, tables flattened (said so).
* DOCX — python-docx (MIT): headings become sections, tables Markdown tables;
  fallback: the repo's stdlib parser ``bossman.file_intel``.
* XLSX — openpyxl (MIT) read-only; fallback ``bossman.file_intel``.
* PPTX / EPUB — ``bossman.file_intel`` (stdlib zip+xml, already in the repo).
* HTML — trafilatura (Apache-2.0) if installed, else BeautifulSoup (MIT),
  else the stdlib HTML parser.
* CSV / TXT / MD — stdlib.
* Images — Pillow metadata; OCR only with ``use_ocr`` and a local RapidOCR.

A missing library is an honest ``UNSUPPORTED`` with the install hint, never a crash.
"""
from __future__ import annotations

import csv
import io
import logging
import re
from dataclasses import asdict, dataclass, field
from html.parser import HTMLParser
from importlib import metadata
from pathlib import Path
from typing import Any

from .config import ensure_repo_imports

PASS, PARTIAL, UNSUPPORTED, FAIL = "PASS", "PARTIAL", "UNSUPPORTED", "FAIL"
STATUSES = (PASS, PARTIAL, UNSUPPORTED, FAIL)

FRAGMENT_CHARS = 6000
ROWS_PER_BLOCK = 200
MAX_SHEET_ROWS = 20_000
MAX_PDF_PAGES = 2000
MAX_TOTAL_CHARS = 5_000_000
MAX_CELL = 300

TABLES_FLATTENED = ("Таблицы PDF извлечены pypdf как плоский текст: столбцы могут быть слиты. "
                    "Для таблиц в Markdown включите Docling (--docling), если он установлен.")


@dataclass
class Unit:
    kind: str      # page | section | fragment | sheet | rows | slide | chapter | meta
    number: int    # page number for PDFs; ordinal otherwise
    label: str     # "Стр. 12", "Раздел 3: Введение", "Лист «Q1», строки 1–200", ...
    text: str


@dataclass
class ParseResult:
    status: str
    parser: str
    units: list[Unit] = field(default_factory=list)
    page_count: int | None = None
    title: str = ""
    reason: str = ""
    warnings: list[str] = field(default_factory=list)
    tables: str = ""             # flattened | markdown | ""
    meta: dict = field(default_factory=dict)

    def text(self) -> str:
        return "\n\n".join(u.text for u in self.units if u.text)

    def to_json(self) -> dict:
        return asdict(self)

    @classmethod
    def from_json(cls, data: dict) -> "ParseResult":
        units = [Unit(**u) for u in data.get("units", [])]
        fields = {k: v for k, v in data.items() if k in cls.__dataclass_fields__ and k != "units"}
        return cls(units=units, **fields)


# --------------------------------------------------------------------------- helpers

def _version(dist: str) -> str:
    try:
        return metadata.version(dist)
    except metadata.PackageNotFoundError:
        return ""


def _named(name: str, dist: str) -> str:
    ver = _version(dist)
    return f"{name} {ver}".strip()


_CTRL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")
_MD_STRUCTURE = re.compile(r"^(\s*)(#{1,6}(?:\s|$)|```|~~~)")


def clean_text(text: str, *, escape_markdown: bool = True) -> str:
    """Normalise extracted text. Lines that would open a Markdown heading or code
    fence are escaped: a stray '#' from a PDF must not cut the page section (and
    its citation) in half when the memory index splits notes by headings."""
    text = _CTRL.sub("", (text or "").replace("\r\n", "\n").replace("\r", "\n"))
    lines = [line.rstrip() for line in text.split("\n")]
    if escape_markdown:
        lines = [_MD_STRUCTURE.sub(lambda m: m.group(1) + "\\" + m.group(2), line) for line in lines]
    return re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()


def decode_text(data: bytes) -> str:
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16", "replace")
    for encoding in ("utf-8-sig", "cp1251"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", "replace")


def split_text(text: str, limit: int) -> list[str]:
    """Split at paragraph, then line, then hard boundaries; every piece <= limit."""
    if len(text) <= limit:
        return [text] if text else []
    pieces: list[str] = []
    buf = ""

    def add(segment: str, sep: str) -> None:
        nonlocal buf
        while len(segment) > limit:
            if buf:
                pieces.append(buf)
                buf = ""
            pieces.append(segment[:limit])
            segment = segment[limit:]
        if buf and len(buf) + len(sep) + len(segment) > limit:
            pieces.append(buf)
            buf = segment
        else:
            buf = f"{buf}{sep}{segment}" if buf else segment

    for para in re.split(r"\n{2,}", text):
        if len(para) <= limit:
            add(para, "\n\n")
        else:
            for i, line in enumerate(para.split("\n")):
                add(line, "\n\n" if i == 0 else "\n")
    if buf:
        pieces.append(buf)
    return [p.strip() for p in pieces if p.strip()]


def markdown_table(rows: list[list[str]]) -> str:
    rows = [[" ".join(str(c).split())[:MAX_CELL].replace("|", "\\|") for c in r] for r in rows if r]
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    out = ["| " + " | ".join(rows[0]) + " |", "|" + "---|" * width]
    out += ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return "\n".join(out)


def _fragments(text: str, *, kind: str = "fragment", label: str = "Фрагмент") -> list[Unit]:
    return [Unit(kind, i, f"{label} {i}", piece)
            for i, piece in enumerate(split_text(text, FRAGMENT_CHARS), start=1)]


def _row_blocks(name: str, rows: list[list[str]], start_number: int) -> list[Unit]:
    """A table as blocks of ROWS_PER_BLOCK rows; the header row repeats in every block."""
    if not rows:
        return []
    header, body = rows[0], rows[1:] or []
    units = []
    number = start_number
    for start in range(0, max(len(body), 1), ROWS_PER_BLOCK):
        chunk = body[start:start + ROWS_PER_BLOCK]
        first, last = start + 2, start + 1 + len(chunk)
        span = f"строки {first}–{last}" if chunk else "только заголовок"
        units.append(Unit("sheet", number, f"{name}, {span}", markdown_table([header] + chunk)))
        number += 1
    return units


def _file_intel():
    ensure_repo_imports()
    from bossman import file_intel  # repo's own stdlib parse-router (bossman-core)
    return file_intel


# --------------------------------------------------------------------------- PDF

def _pdf_docling(path: Path, max_pages: int) -> ParseResult:
    ensure_repo_imports()
    from bcc.oss import docling as repo_docling  # native pipeline: no model download, no remote services

    result = repo_docling._convert(path.read_bytes(), path.name, max_pages)
    status = str(getattr(result.status, "value", result.status)).lower()
    if status != "success":
        raise RuntimeError(f"docling status {status}")
    doc = result.document
    numbers = sorted(int(n) for n in (getattr(doc, "pages", {}) or {}))
    units = [Unit("page", n, f"Стр. {n}", clean_text(doc.export_to_markdown(page_no=n), escape_markdown=False))
             for n in numbers]
    if not any(u.text for u in units):
        raise RuntimeError("docling produced no text")
    return ParseResult(PASS, _named("docling", "docling"), units, page_count=len(numbers), tables="markdown")


class _CountWarnings(logging.Handler):
    """pypdf logs per-font warnings (e.g. missing fontTools) to stderr: dozens per report.
    Counted and summarised once in the manifest instead of flooding the console."""

    def __init__(self) -> None:
        super().__init__(logging.WARNING)
        self.count = 0
        self.fonttools = False

    def emit(self, record: logging.LogRecord) -> None:
        self.count += 1
        self.fonttools = self.fonttools or "fontTools" in record.getMessage()


def parse_pdf(path: Path, *, max_pages: int = MAX_PDF_PAGES, use_docling: bool = False) -> ParseResult:
    logger = logging.getLogger("pypdf")
    counter, propagate = _CountWarnings(), logger.propagate
    logger.addHandler(counter)
    logger.propagate = False
    try:
        result = _parse_pdf(path, max_pages=max_pages, use_docling=use_docling)
    finally:
        logger.removeHandler(counter)
        logger.propagate = propagate
    if counter.count:
        hint = " (нет пакета fontTools: кодировка части шрифтов может быть неточной)" if counter.fonttools else ""
        result.warnings.append(f"pypdf: {counter.count} предупреждений при извлечении{hint}")
    return result


def _parse_pdf(path: Path, *, max_pages: int, use_docling: bool) -> ParseResult:
    warnings: list[str] = []
    if use_docling:
        try:
            return _pdf_docling(path, max_pages)
        except ImportError:
            warnings.append("Docling не установлен: использован pypdf")
        except Exception as exc:  # noqa: BLE001 - any Docling failure falls back to pypdf, said aloud
            warnings.append(f"Docling не справился ({type(exc).__name__}): использован pypdf")
    try:
        from pypdf import PdfReader
    except ImportError:
        return ParseResult(UNSUPPORTED, "none", reason="PDF: не установлен pypdf (pip install pypdf)",
                           warnings=warnings)
    parser = _named("pypdf", "pypdf")
    try:
        reader = PdfReader(str(path), strict=False)
        if reader.is_encrypted:
            try:
                opened = reader.decrypt("")
            except Exception:  # noqa: BLE001 - unknown crypto handler == cannot open
                opened = 0
            if not opened:
                return ParseResult(FAIL, parser, reason="PDF защищён паролем", warnings=warnings)
        total = len(reader.pages)
    except Exception as exc:  # noqa: BLE001 - a damaged PDF is a FAIL with a reason, not a crash
        return ParseResult(FAIL, parser, reason=f"PDF не читается: {type(exc).__name__}", warnings=warnings)
    title = ""
    try:
        title = str((reader.metadata or {}).get("/Title") or "").strip()
    except Exception:  # noqa: BLE001 - metadata is optional
        title = ""
    units: list[Unit] = []
    empty: list[int] = []
    errors: list[int] = []
    for index in range(min(total, max_pages)):
        number = index + 1
        try:
            raw = reader.pages[index].extract_text() or ""
        except Exception:  # noqa: BLE001 - one bad page must not lose the other pages
            raw = ""
            errors.append(number)
        text = clean_text(raw)
        if not text:
            empty.append(number)
        units.append(Unit("page", number, f"Стр. {number}", text))
    with_text = len(units) - len(empty)
    meta = {"empty_pages": empty, "error_pages": errors}
    warnings.append(TABLES_FLATTENED)
    if with_text == 0:
        return ParseResult(FAIL, parser, units, page_count=total, title=title, meta=meta, warnings=warnings,
                           reason="в PDF нет текстового слоя (скан?): локальный OCR не установлен")
    reasons = []
    if total > max_pages:
        reasons.append(f"разобраны первые {max_pages} из {total} стр.")
    if errors:
        reasons.append(f"ошибка извлечения на стр. {', '.join(map(str, errors[:20]))}")
    if len(empty) > max(2, total // 10):
        reasons.append(f"без текста {len(empty)} стр. из {total} (сканы?)")
    return ParseResult(PARTIAL if reasons else PASS, parser, units, page_count=total, title=title,
                       reason="; ".join(reasons), warnings=warnings, tables="flattened", meta=meta)


# --------------------------------------------------------------------------- DOCX

_HEADING_STYLE = re.compile(r"^(heading|заголовок|title|название)\b", re.I)


def _docx_python_docx(path: Path) -> ParseResult:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = docx.Document(str(path))
    try:
        title = (document.core_properties.title or "").strip()
    except Exception:  # noqa: BLE001
        title = ""
    blocks = document.iter_inner_content() if hasattr(document, "iter_inner_content") else document.paragraphs
    sections: list[tuple[str, list[str]]] = []
    heading, lines, tables = "", [], 0

    def flush():
        if heading or lines:
            sections.append((heading, list(lines)))

    for block in blocks:
        if isinstance(block, Paragraph):
            text = " ".join(block.text.split())
            if not text:
                continue
            try:
                style = block.style.name if block.style is not None else ""
            except Exception:  # noqa: BLE001 - broken style references in foreign files
                style = ""
            if _HEADING_STYLE.match(style or ""):
                flush()
                heading, lines = text, []
            else:
                lines.append(("- " + text) if (style or "").lower().startswith(("list", "спис")) else text)
        elif isinstance(block, Table):
            rows = [[cell.text for cell in row.cells] for row in block.rows]
            table = markdown_table(rows)
            if table:
                tables += 1
                lines.append(table)
    flush()
    units: list[Unit] = []
    if len(sections) == 1 and not sections[0][0]:
        units = _fragments(clean_text("\n\n".join(sections[0][1]), escape_markdown=False))
    else:
        for number, (head, body) in enumerate(sections, start=1):
            label = f"Раздел {number}: {head[:80]}" if head else f"Раздел {number}"
            units.append(Unit("section", number, label, clean_text("\n\n".join(body), escape_markdown=False)))
    return ParseResult(PASS, _named("python-docx", "python-docx"), units, title=title,
                       tables="markdown" if tables else "", meta={"tables": tables})


def parse_docx(path: Path) -> ParseResult:
    warnings: list[str] = []
    try:
        import docx  # noqa: F401
        try:
            return _docx_python_docx(path)
        except Exception as exc:  # noqa: BLE001 - fall back to the stdlib parser
            warnings.append(f"python-docx не справился ({type(exc).__name__}): использован stdlib-разбор")
    except ImportError:
        warnings.append("python-docx не установлен: использован stdlib-разбор (без заголовков)")
    fi = _file_intel()
    sections = fi._parse_docx(path.read_bytes())
    text_parts = [s.text for s in sections if s.kind == "text" and s.text]
    tables = [markdown_table([list(r) for r in s.table]) for s in sections if s.kind == "table"]
    body = "\n\n".join(text_parts + tables)
    if tables:
        warnings.append("таблицы документа перенесены в конец текста")
    units = _fragments(clean_text(body, escape_markdown=False))
    return ParseResult(PASS if units else FAIL, "bossman.file_intel (stdlib)", units,
                       reason="" if units else "в DOCX нет текста", warnings=warnings,
                       tables="markdown" if tables else "")


# --------------------------------------------------------------------------- XLSX / CSV

def parse_xlsx(path: Path) -> ParseResult:
    units: list[Unit] = []
    warnings: list[str] = []
    try:
        import openpyxl
    except ImportError:
        openpyxl = None
    if openpyxl is not None:
        workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        try:
            for sheet in workbook.worksheets:
                rows: list[list[str]] = []
                for row in sheet.iter_rows(values_only=True):
                    if len(rows) >= MAX_SHEET_ROWS:
                        warnings.append(f"лист «{sheet.title}»: взяты первые {MAX_SHEET_ROWS} строк")
                        break
                    values = ["" if v is None else str(v) for v in row]
                    while values and not values[-1].strip():
                        values.pop()
                    if any(v.strip() for v in values):
                        rows.append(values)
                units += _row_blocks(f"Лист «{sheet.title}»", rows, len(units) + 1)
        finally:
            workbook.close()
        parser = _named("openpyxl", "openpyxl")
    else:
        fi = _file_intel()
        for section in fi._parse_xlsx(path.read_bytes()):
            name = (section.meta or {}).get("sheet") or section.ref
            units += _row_blocks(f"Лист «{name}»", [list(r) for r in section.table], len(units) + 1)
        parser = "bossman.file_intel (stdlib)"
        warnings.append("openpyxl не установлен: использован stdlib-разбор")
    status = PARTIAL if any("первые" in w for w in warnings) else PASS
    if not units:
        return ParseResult(FAIL, parser, reason="в таблице нет данных", warnings=warnings)
    return ParseResult(status, parser, units, warnings=warnings, tables="markdown",
                       reason="; ".join(w for w in warnings if "первые" in w))


def parse_csv(path: Path) -> ParseResult:
    text = decode_text(path.read_bytes())
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel_tab if path.suffix.lower() == ".tsv" else csv.excel
    rows = [row for row in csv.reader(io.StringIO(text), dialect) if any(c.strip() for c in row)]
    truncated = len(rows) > MAX_SHEET_ROWS
    rows = rows[:MAX_SHEET_ROWS]
    units = _row_blocks("Таблица", rows, 1)
    for unit in units:
        unit.kind = "rows"
    if not units:
        return ParseResult(FAIL, "csv (stdlib)", reason="в CSV нет строк")
    return ParseResult(PARTIAL if truncated else PASS, "csv (stdlib)", units, tables="markdown",
                       reason=f"взяты первые {MAX_SHEET_ROWS} строк" if truncated else "")


# --------------------------------------------------------------------------- text / markdown / html

_MD_HEADING = re.compile(r"(?m)^(#{1,3})\s+(.+?)\s*#*\s*$")


def _markdown_sections(text: str) -> list[Unit]:
    matches = list(_MD_HEADING.finditer(text))
    if not matches:
        return _fragments(clean_text(text, escape_markdown=False))
    units: list[Unit] = []
    intro = text[:matches[0].start()].strip()
    if intro:
        units.append(Unit("section", 1, "Раздел 1: начало", clean_text(intro, escape_markdown=False)))
    for i, match in enumerate(matches):
        end = matches[i + 1].start() if i + 1 < len(matches) else len(text)
        body = clean_text(text[match.end():end], escape_markdown=False)
        number = len(units) + 1
        units.append(Unit("section", number, f"Раздел {number}: {match.group(2)[:80]}", body))
    return units


def parse_text(path: Path, kind: str) -> ParseResult:
    text = decode_text(path.read_bytes())
    if kind == "md":
        units = _markdown_sections(text)
        title = next((m.group(2) for m in _MD_HEADING.finditer(text) if m.group(1) == "#"), "")
    else:
        units = _fragments(clean_text(text))
        title = ""
    if not any(u.text for u in units):
        return ParseResult(FAIL, "text (stdlib)", reason="файл не содержит текста")
    return ParseResult(PASS, "text (stdlib)", units, title=title)


class _HtmlText(HTMLParser):
    _BLOCK = {"p", "div", "br", "li", "tr", "section", "article", "blockquote", "pre", "table", "ul", "ol"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self.skip = 0
        self.title = ""
        self._in_title = False

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style", "noscript", "template"):
            self.skip += 1
        elif tag == "title":
            self._in_title = True
        elif tag in ("h1", "h2", "h3"):
            self.parts.append(f"\n\n{'#' * int(tag[1])} ")
        elif tag in self._BLOCK:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in ("script", "style", "noscript", "template") and self.skip:
            self.skip -= 1
        elif tag == "title":
            self._in_title = False
        elif tag in ("h1", "h2", "h3") or tag in self._BLOCK:
            self.parts.append("\n")

    def handle_data(self, data):
        if self._in_title:
            self.title += data
        elif not self.skip:
            self.parts.append(data)


def parse_html(path: Path) -> ParseResult:
    raw = decode_text(path.read_bytes())
    markdown, title, parser = "", "", ""
    try:
        import trafilatura

        markdown = trafilatura.extract(raw, output_format="markdown", include_tables=True,
                                       include_comments=False) or ""
        parser = _named("trafilatura", "trafilatura")
    except ImportError:
        pass
    if not markdown:
        try:
            from bs4 import BeautifulSoup

            soup = BeautifulSoup(raw, "html.parser")
            title = (soup.title.get_text(" ", strip=True) if soup.title else "")
            for tag in soup(["script", "style", "noscript", "template", "head"]):
                tag.decompose()
            for level in ("h1", "h2", "h3"):
                for tag in soup.find_all(level):
                    tag.replace_with(f"\n\n{'#' * int(level[1])} {tag.get_text(' ', strip=True)}\n")
            markdown = soup.get_text("\n")
            parser = _named("beautifulsoup4", "beautifulsoup4")
        except ImportError:
            reader = _HtmlText()
            reader.feed(raw)
            markdown, title, parser = "".join(reader.parts), reader.title.strip(), "html.parser (stdlib)"
    markdown = "\n".join(" ".join(line.split()) if not line.lstrip().startswith("#") else line.strip()
                         for line in markdown.split("\n"))
    units = _markdown_sections(markdown)
    if not any(u.text for u in units):
        return ParseResult(FAIL, parser, reason="в HTML нет текста")
    return ParseResult(PASS, parser, units, title=" ".join(title.split()))


# --------------------------------------------------------------------------- PPTX / EPUB / images / legacy

def parse_pptx(path: Path) -> ParseResult:
    fi = _file_intel()
    units = [Unit("slide", int((s.meta or {}).get("slide") or i), f"Слайд {(s.meta or {}).get('slide') or i}",
                  clean_text(s.text)) for i, s in enumerate(fi._parse_pptx(path.read_bytes()), start=1)]
    if not any(u.text for u in units):
        return ParseResult(FAIL, "bossman.file_intel (stdlib)", units, reason="на слайдах нет текста")
    return ParseResult(PASS, "bossman.file_intel (stdlib)", units, page_count=len(units))


def parse_epub(path: Path) -> ParseResult:
    fi = _file_intel()
    sections = fi._parse_epub(path.read_bytes())
    title = next((s.text for s in sections if s.ref == "meta=title"), "")
    chapters = [s for s in sections if s.ref.startswith("chapter=")]
    units = [Unit("chapter", i, f"Глава {i}", clean_text(s.text)) for i, s in enumerate(chapters, start=1)]
    if not units:
        return ParseResult(FAIL, "bossman.file_intel (stdlib)", reason="в EPUB не найдено глав")
    return ParseResult(PASS, "bossman.file_intel (stdlib)", units, title=title)


def parse_image(path: Path, *, use_ocr: bool = False) -> ParseResult:
    meta: dict[str, Any] = {"bytes": path.stat().st_size, "format": path.suffix.lstrip(".").lower()}
    parser = "metadata"
    try:
        from PIL import Image

        with Image.open(path) as image:
            meta.update(width=image.width, height=image.height, format=image.format or meta["format"],
                        mode=image.mode)
            try:
                exif = image.getexif()
                shot = exif.get(36867) or exif.get(306)
                if shot:
                    meta["taken"] = str(shot)
            except Exception:  # noqa: BLE001 - EXIF is optional
                pass
        parser = _named("Pillow", "pillow")
    except ImportError:
        pass
    except Exception as exc:  # noqa: BLE001 - unreadable image: keep the file, say why
        meta["error"] = type(exc).__name__
    size = f"{meta.get('width', '?')}×{meta.get('height', '?')}"
    desc = f"Изображение {meta.get('format')} {size}, {meta['bytes'] / 1024:.0f} КБ"
    if meta.get("taken"):
        desc += f", снято {meta['taken']}"
    text, reason, status = "", "", UNSUPPORTED
    if use_ocr:
        try:
            from rapidocr_onnxruntime import RapidOCR

            found, _ = RapidOCR()(str(path))
            text = "\n".join(item[1] for item in (found or []) if len(item) > 1)
            parser += " + rapidocr"
            status = PASS if text.strip() else PARTIAL
            reason = "" if text.strip() else "OCR не нашёл текста"
        except ImportError:
            reason = "OCR включён, но RapidOCR не установлен: сохранены файл и метаданные"
    else:
        reason = "текст изображения не распознан (OCR выключен/не установлен): сохранены файл и метаданные"
    body = desc + ("\n\nРаспознанный текст:\n\n" + clean_text(text) if text.strip() else "")
    return ParseResult(status, parser, [Unit("meta", 1, "Описание", body)], reason=reason, meta=meta)


def parse(path: Path, kind: str, *, max_pages: int = MAX_PDF_PAGES, use_docling: bool = False,
          use_ocr: bool = False) -> ParseResult:
    """Parse one local file. Never raises: an unexpected error becomes FAIL with the error type."""
    path = Path(path)
    try:
        if kind == "pdf":
            result = parse_pdf(path, max_pages=max_pages, use_docling=use_docling)
        elif kind == "docx":
            result = parse_docx(path)
        elif kind == "xlsx":
            result = parse_xlsx(path)
        elif kind == "csv":
            result = parse_csv(path)
        elif kind in ("txt", "md"):
            result = parse_text(path, kind)
        elif kind == "html":
            result = parse_html(path)
        elif kind == "pptx":
            result = parse_pptx(path)
        elif kind == "epub":
            result = parse_epub(path)
        elif kind == "image":
            result = parse_image(path, use_ocr=use_ocr)
        elif kind == "legacy":
            result = ParseResult(UNSUPPORTED, "none",
                                 reason=f"нет локального парсера для {path.suffix.lower()}: файл сохранён в библиотеке; "
                                        "пересохраните в DOCX/XLSX/PDF и выполните reparse")
        else:
            result = ParseResult(UNSUPPORTED, "none", reason=f"тип {kind or path.suffix} не поддерживается")
    except Exception as exc:  # noqa: BLE001 - one broken document never stops the library
        return ParseResult(FAIL, "error", reason=f"разбор упал: {type(exc).__name__}: {str(exc)[:160]}")
    total = 0
    for i, unit in enumerate(result.units):
        total += len(unit.text)
        if total > MAX_TOTAL_CHARS:
            result.units = result.units[:i]
            result.status = PARTIAL if result.status == PASS else result.status
            result.reason = "; ".join(x for x in (result.reason, f"текст обрезан на {MAX_TOTAL_CHARS} знаков") if x)
            break
    return result
